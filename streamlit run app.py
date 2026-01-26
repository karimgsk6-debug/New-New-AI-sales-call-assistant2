import streamlit as st
from PIL import Image
import requests
from io import BytesIO
from datetime import datetime
from groq import Groq

# =========================
# OPTIONAL WORD DOWNLOAD
# =========================
try:
    from docx import Document
    from io import BytesIO as io_bytes
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

# =========================
# GROQ CLIENT
# =========================
client = Groq(api_key="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")

# =========================
# BRAND MASTER DATA
# =========================
brand_data = {
    "shingrix": {
        "display":"Shingrix",
        "segments":["R – Reach","A – Acquisition","C – Conversion","E – Engagement"],
        "personas":["Uncommitted Vaccinator","Reluctant Efficiency","Patient Influenced","Committed Vaccinator"],
        "barriers":["HCP does not consider HZ a risk","No time for discussion","Cost concerns","Not convinced of efficacy"],
        "specialties":["GP","Dermatologist","Geriatrician"],
        "references_path":".devcontainer/references/shingrix/",
        "sales_path":".devcontainer/SalesModule/shingrix/",
        "call_flow":["Prepare","Engage","Create Opportunities","Influence","Impact GSO","Post-call Analysis"],
        "leaflet":"https://example.com/shingrix-leaflet",
        "image":"https://www.oma-apteekki.fi/WebRoot/NA/Shops/na/67D6/48DA/D0B0/D959/ECAF/0A3C/0E02/D573/3ad67c4e-e1fb-4476-a8a0-873423d8db42_3Dimage.png"
    },
    "jemperli": {
        "display":"Jemperli",
        "segments":["Target Identification","Trial Adoption","Routine Use","Advocacy"],
        "personas":["Data-Driven Oncologist","Skeptical Specialist","Innovator Prescriber","Late Adopter"],
        "barriers":["Unfamiliar with immunotherapy","Safety concerns","Limited eligibility","Access/reimbursement issues"],
        "specialties":["Oncologist","Medical Oncologist"],
        "references_path":".devcontainer/references/jemperli/",
        "sales_path":".devcontainer/SalesModule/jemperli/",
        "call_flow":["COCO","Anchor","Engage","Close"],
        "leaflet":"https://example.com/jemperli-leaflet",
        "image":"https://via.placeholder.com/300x150.png?text=Jemperli"
    },
    "trelegy": {
        "display":"Trelegy",
        "segments":["Awareness","Diagnosis","Adoption","Adherence"],
        "personas":["Primary Care COPD Prescriber","Pulmonologist","Respiratory Nurse"],
        "barriers":["Formulary access","Inhaler technique","Side effect concerns","Cost/coverage"],
        "specialties":["GP","Pulmonologist","Respiratory Specialist"],
        "references_path":".devcontainer/references/trelegy/",
        "sales_path":".devcontainer/SalesModule/trelegy/",
        "call_flow":["Prepare","Engage","Demonstrate","Address Access","Close"],
        "leaflet":"https://example.com/trelegy-leaflet",
        "image":"https://www.example.com/trelegy.png"
    }
}

# =========================
# SESSION STATE
# =========================
if "chat_history" not in st.session_state:
    st.session_state.chat_history = []

# =========================
# LANGUAGE
# =========================
language = st.radio("Select Language / اختر اللغة", ["English", "العربية"])

# =========================
# HEADER
# =========================
st.title("🧠 AI Sales Call Assistant")

# =========================
# SIDEBAR – BRAND DRIVEN
# =========================
st.sidebar.header("Filters & Options")

brand_key = st.sidebar.selectbox(
    "Select Brand",
    options=list(brand_data.keys()),
    format_func=lambda x: brand_data[x]["display"]
)

brand_cfg = brand_data[brand_key]

segment = st.sidebar.selectbox("Segment", brand_cfg["segments"])
persona = st.sidebar.selectbox("Persona", brand_cfg["personas"])
barrier = st.sidebar.multiselect("Doctor Barriers", brand_cfg["barriers"])
specialty = st.sidebar.selectbox("Specialty", brand_cfg["specialties"])

objective = st.sidebar.selectbox("Objective", ["Awareness", "Adoption", "Retention"])
response_length = st.sidebar.selectbox("Response Length", ["Short", "Medium", "Long"])
response_tone = st.sidebar.selectbox("Response Tone", ["Formal", "Casual", "Friendly", "Persuasive"])

# =========================
# CLEAR CHAT
# =========================
if st.button("🗑️ Clear Chat"):
    st.session_state.chat_history = []

# =========================
# CHAT DISPLAY
# =========================
def display_chat():
    for msg in st.session_state.chat_history:
        if msg["role"] == "user":
            st.markdown(f"**🧑 You:** {msg['content']}")
        else:
            st.markdown(f"**🤖 AI:** {msg['content']}")

display_chat()

# =========================
# CHAT INPUT
# =========================
with st.form("chat_form", clear_on_submit=True):
    user_input = st.text_input("Type your message...")
    submitted = st.form_submit_button("Send")

# =========================
# AI LOGIC
# =========================
if submitted and user_input.strip():

    st.session_state.chat_history.append({
        "role":"user",
        "content":user_input,
        "time":datetime.now().strftime("%H:%M")
    })

    flow_str = " → ".join(brand_cfg["call_flow"])

    prompt = f"""
Language: {language}
User Request: {user_input}

Brand: {brand_cfg['display']}
Segment: {segment}
Persona: {persona}
Specialty: {specialty}
Doctor Barriers: {', '.join(barrier) if barrier else 'None'}
Objective: {objective}

Sales Call Flow:
{flow_str}

Use APACT (Acknowledge → Probing → Answer → Confirm → Transition).
Provide compliant, practical sales guidance.
Tone: {response_tone}
Length: {response_length}
"""

    response = client.chat.completions.create(
        model="llama-3.3-70b-versatile",
        messages=[
            {"role":"system","content":f"You are a pharmaceutical sales excellence AI responding in {language}."},
            {"role":"user","content":prompt}
        ],
        temperature=0.7
    )

    ai_output = response.choices[0].message.content

    st.session_state.chat_history.append({
        "role":"ai",
        "content":ai_output,
        "time":datetime.now().strftime("%H:%M")
    })

    display_chat()

# =========================
# WORD DOWNLOAD
# =========================
if DOCX_AVAILABLE:
    ai_msgs = [m["content"] for m in st.session_state.chat_history if m["role"] == "ai"]
    if ai_msgs:
        doc = Document()
        doc.add_heading("AI Sales Call Output", 0)
        doc.add_paragraph(ai_msgs[-1])
        buffer = io_bytes()
        doc.save(buffer)
        st.download_button("📥 Download Word", buffer.getvalue(), file_name="AI_Sales_Call.docx")

# =========================
# BRAND LEAFLET
# =========================
st.markdown(f"[📄 Brand Leaflet – {brand_cfg['display']}]({brand_cfg['leaflet']})")
# -------------------------
# Footer disclaimer
# -------------------------
st.markdown("""
<div class="fixed-disclaimer">
💡 This tool is for internal sales support purposes only. All medical info should be verified from official sources. 
</div>
""",unsafe_allow_html=True)
